# ============================================================
# RESUME TAILORING AND DOCX EXPORT
#
# data/resume_template.md          committed layout with <KEY> placeholders
# .env                             personal details (NAME, EMPLOYER_1, ...)
# data/private/experience.yaml     master content (gitignored)
# data/private/resumes/<slug>.yaml tailored copy per job (gitignored)
# data/output/<slug>.docx          generated resume (gitignored)
#
# Usage (from repo root):
#   python src/resume.py general
#   python src/resume.py list
#   python src/resume.py tailor "Mannai Trading" --role "Systems Engineer" \
#       --bullets w22_01,w22_04,w12_01 --projects project_2,project_4
#   python src/resume.py render "Mannai Trading"
# ============================================================

import argparse
import os
import re
from datetime import datetime

import yaml
from docx import Document
from dotenv import dotenv_values

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
TEMPLATE = os.path.join(ROOT, "data", "resume_template.md")
MASTER = os.path.join(ROOT, "data", "private", "experience.yaml")
TAILORED_DIR = os.path.join(ROOT, "data", "private", "resumes")
OUTPUT_DIR = os.path.join(ROOT, "data", "output")

EXPERIENCE_PREFIX = "WORK_EXPERIENCE_"


def slugify(company):
    return re.sub(r"[^a-z0-9]+", "-", company.lower()).strip("-")


def load_yaml(path):
    with open(path, "r", encoding="utf-8") as file:
        return yaml.safe_load(file)


def experience_keys(data):
    return [key for key in data if key.startswith(EXPERIENCE_PREFIX)]


# ============================================================
# LIST — show every bullet and project id in the master file
# ============================================================

def list_master():
    master = load_yaml(MASTER)

    for key in experience_keys(master):
        print(f"\n{key}")
        for bullet in master[key]:
            print(f"  {bullet['id']}: {bullet['text']}")

    print("\nprojects")
    for project in master.get("projects", []):
        print(f"  {project['id']}: {project['name']}")


# ============================================================
# TAILOR — copy selected bullets and projects into a job file
# ============================================================

def tailor(company, role, bullet_ids, project_ids):
    master = load_yaml(MASTER)

    known_bullets = {b["id"] for key in experience_keys(master) for b in master[key]}
    known_projects = {p["id"] for p in master.get("projects", [])}
    unknown = [i for i in bullet_ids if i not in known_bullets]
    unknown += [i for i in project_ids if i not in known_projects]
    if unknown:
        raise SystemExit(f"Unknown ids (not in master): {', '.join(unknown)}")

    tailored = {
        "company": company,
        "role": role,
        "created": datetime.now().strftime("%Y-%m-%d"),
        "professional_summary": master["professional_summary"],
        "technical_skills": master["technical_skills"],
    }

    # Keep master order, only the selected bullets
    for key in experience_keys(master):
        tailored[key] = [b for b in master[key] if b["id"] in bullet_ids]

    tailored["projects"] = [p for p in master.get("projects", []) if p["id"] in project_ids]

    os.makedirs(TAILORED_DIR, exist_ok=True)
    path = os.path.join(TAILORED_DIR, f"{slugify(company)}.yaml")
    with open(path, "w", encoding="utf-8") as file:
        yaml.safe_dump(tailored, file, sort_keys=False, allow_unicode=True, width=10000)

    print(f"Tailored resume saved: {os.path.relpath(path, ROOT)}")


# ============================================================
# RENDER — fill template and write .docx
# ============================================================

def build_markdown(tailored):
    env = dotenv_values(os.path.join(ROOT, ".env"))

    blocks = {
        "PROFESSIONAL_SUMMARY": tailored["professional_summary"],
        "TECHNICAL_SKILLS": "\n".join(
            f"- **{label}:** {skills}" for label, skills in tailored["technical_skills"].items()
        ),
        "PROJECTS": "\n\n".join(
            f"### {p['name']}\n\n{p['description']}\n\n*Technology context: {p['tech_stack']}*"
            for p in tailored.get("projects", [])
        ),
    }
    for key in experience_keys(tailored):
        blocks[key] = "\n".join(f"- {b['text']}" for b in tailored[key])

    with open(TEMPLATE, "r", encoding="utf-8") as file:
        text = file.read()

    def replace(match):
        key = match.group(1)
        if key in blocks:
            return blocks[key]
        value = env.get(key)
        if value:
            return value
        raise SystemExit(f"No value for placeholder <{key}> (add it to .env or the tailored file)")

    text = re.sub(r"<([A-Z][A-Z0-9_]*)>", replace, text)

    # Drop roles with no selected bullets: heading, date line and empty block
    return re.sub(r"### [^\n]*\n\n\*[^\n]*\*\n\n(?=\n|###|## )", "", text)


def add_runs(paragraph, text):
    # Supports **bold** and *italic* from the template only
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and len(part) > 1:
            paragraph.add_run(part[1:-1]).italic = True
        elif part:
            paragraph.add_run(part)


def markdown_to_docx(markdown, path):
    doc = Document()

    for line in markdown.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            add_runs(doc.add_paragraph(), line)

    doc.save(path)


def render(company):
    path = os.path.join(TAILORED_DIR, f"{slugify(company)}.yaml")
    if not os.path.exists(path):
        raise SystemExit(f"No tailored resume for '{company}'. Run tailor first.")

    save_docx(load_yaml(path), slugify(company))


def render_general():
    # Full resume straight from the master file, no tailoring
    save_docx(load_yaml(MASTER), "general")


def save_docx(data, name):
    markdown = build_markdown(data)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    output = os.path.join(OUTPUT_DIR, f"{name}.docx")
    markdown_to_docx(markdown, output)
    print(f"Resume saved: {os.path.relpath(output, ROOT)}")


# ============================================================
# CLI
# ============================================================

def split_ids(value):
    return [i.strip() for i in value.split(",") if i.strip()] if value else []


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Tailor and export resumes")
    commands = parser.add_subparsers(dest="command", required=True)

    commands.add_parser("list", help="Show bullet and project ids in the master file")
    commands.add_parser("general", help="Write the full resume from the master file")

    tailor_cmd = commands.add_parser("tailor", help="Create a tailored resume for a job")
    tailor_cmd.add_argument("company")
    tailor_cmd.add_argument("--role", default="")
    tailor_cmd.add_argument("--bullets", required=True, help="Comma-separated bullet ids")
    tailor_cmd.add_argument("--projects", default="", help="Comma-separated project ids")

    render_cmd = commands.add_parser("render", help="Write the .docx for a tailored resume")
    render_cmd.add_argument("company")

    args = parser.parse_args()

    if args.command == "list":
        list_master()
    elif args.command == "general":
        render_general()
    elif args.command == "tailor":
        tailor(args.company, args.role, split_ids(args.bullets), split_ids(args.projects))
    elif args.command == "render":
        render(args.company)
